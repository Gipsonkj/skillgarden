"""Download permitted skills from _research/*.json, zip per topic, write the xlsx index.
Run: uv run --with openpyxl python _tools/build.py [--no-download]"""
import json, os, re, shutil, subprocess, sys, tarfile, glob, zipfile
from pathlib import Path

# The repo holds _tools and _research; the library (skills/, zips/, _cache, the xlsx) sits next to it.
APP = Path(__file__).resolve().parent.parent
ROOT = Path(os.environ.get("SKILLGARDEN_LIBRARY", APP.parent)).resolve()
RES, SK, ZP = APP / "_research", ROOT / "skills", ROOT / "zips"
CACHE = Path(os.environ.get("SG_CACHE", ROOT / "_cache"))
ORDER = ["ai_agents", "website_building", "poster_design", "motion_animation", "trading_finance",
         "coding_practices", "ai_video", "audio_generation", "image_creation", "storyboarding",
         "social_media", "content_creation", "app_building", "cloud_devops", "backend_databases",
         "automation", "ad_creation", "google_ads", "linkedin_automation", "instagram_automation",
         "frontend_ui_design", "testing_qa", "security", "data_analysis", "seo", "docs_office",
         "research_science", "product_management", "claude_meta", "figma_design",
         "token_efficiency", "presentations", "open_models", "course_design", "email_marketing",
         "ecommerce", "career", "3d_modeling"]

def safe(s):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s).strip("_")[:60]

def fetch_repo(repo, branch):
    dest = CACHE / f"{repo.replace('/', '__')}@{safe(branch)}"
    if dest.exists():
        return dest
    tgz = CACHE / (dest.name + ".tar.gz")
    CACHE.mkdir(parents=True, exist_ok=True)
    with open(tgz, "wb") as f:
        r = subprocess.run(["gh", "api", f"repos/{repo}/tarball/{branch}"], stdout=f, stderr=subprocess.PIPE)
    if r.returncode:
        tgz.unlink(missing_ok=True)
        raise RuntimeError(r.stderr.decode()[:200])
    tmp = CACHE / (dest.name + ".tmp")
    shutil.rmtree(tmp, ignore_errors=True)
    with tarfile.open(tgz) as t:
        t.extractall(tmp, filter="data")
    top = next(tmp.iterdir())
    top.rename(dest)
    shutil.rmtree(tmp)
    tgz.unlink()
    return dest

def license_file(repo_dir):
    for n in ("LICENSE", "LICENSE.md", "LICENSE.txt", "LICENCE", "COPYING"):
        if (repo_dir / n).is_file():
            return repo_dir / n

def main(download=True, only=()):
    topics = []
    for slug in ORDER + sorted({Path(p).stem for p in glob.glob(str(RES / "*.json"))} - set(ORDER)):
        p = RES / f"{slug}.json"
        if p.exists():
            topics.append(json.loads(p.read_text()))
    log = []
    for i, t in enumerate(topics, 1):
        t["folder"] = f"{i:02d}_{safe(t['topic'])}"
        tdir = SK / t["folder"]
        fetch = download and (not only or t["slug"] in only)
        if fetch:
            shutil.rmtree(tdir, ignore_errors=True)
        t["skills"].sort(key=lambda s: s.get("rank", 999))
        for j, s in enumerate(t["skills"], 1):
            s["rank"] = j
            s["status"] = "link only (license)" if not s.get("license_ok") else "pending"
            target = tdir / f"{j:02d}_{safe(s['name'])}"
            if not s.get("license_ok"):
                continue
            if not fetch:
                if (target / "SKILL.md").is_file():
                    s["status"], s["local"] = "downloaded", str(target.relative_to(ROOT))
                continue
            try:
                try:
                    rd = fetch_repo(s["repo"], s.get("branch") or "main")
                except RuntimeError:
                    s["branch"] = subprocess.check_output(["gh", "api", f"repos/{s['repo']}", "--jq", ".default_branch"], text=True).strip()
                    rd = fetch_repo(s["repo"], s["branch"])
                src = rd / s.get("skill_path", "").strip("/")
                if not (src / "SKILL.md").is_file():
                    raise RuntimeError("SKILL.md not found at path")
                shutil.rmtree(target, ignore_errors=True)
                skip = [".git", "node_modules", "__pycache__", ".DS_Store"]
                if src == rd:  # whole-repo skill: drop repo plumbing, keep what SKILL.md uses
                    skip += [".github", "tests", "test"]
                shutil.copytree(src, target, ignore=shutil.ignore_patterns(*skip))
                lf = license_file(rd)
                if lf and not any(target.glob("LICEN*")):
                    shutil.copy(lf, target / "LICENSE.source-repo")
                s["status"] = "downloaded"
                s["local"] = str(target.relative_to(ROOT))
            except Exception as e:
                s["status"] = f"failed: {e}"
                log.append(f"{t['topic']} / {s['name']}: {e}")
        if fetch and tdir.exists():
            ZP.mkdir(exist_ok=True)
            zpath = ZP / f"{t['folder']}.zip"
            with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
                for f in sorted(tdir.rglob("*")):
                    if f.is_file() and not f.is_symlink():
                        z.write(f, f.relative_to(SK))
    (RES / "_build_log.txt").write_text("\n".join(log))
    write_xlsx(topics)
    print(f"{len(topics)} topics; failures: {len(log)}")
    for l in log:
        print("  ", l)

def write_xlsx(topics):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ov = wb.active
    ov.title = "Overview"
    hdr_fill = PatternFill("solid", fgColor="1F3A5F")
    hdr_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    ov.append(["SkillGarden: best public Claude Agent Skills by topic"])
    ov["A1"].font = Font(bold=True, size=14)
    ov.append(["Install: unzip a topic zip and copy any skill folder into ~/.claude/skills/ (personal) or <project>/.claude/skills/ (project). Claude loads it automatically when the task matches its description."])
    ov.append(["Ranking: GitHub stars of the source repo first, then user/expert feedback (skills.sh installs, awesome-list listings, official vendor/Anthropic origin)."])
    ov.append([])
    ov.append(["#", "Topic (tab)", "Skills ranked", "Downloaded", "Link only", "Zip file", "Top 3"])
    for c in ov[5]:
        c.fill, c.font = hdr_fill, hdr_font
    for i, t in enumerate(topics, 1):
        sk = t["skills"]
        ov.append([i, t["topic"], len(sk), sum(s["status"] == "downloaded" for s in sk),
                   sum(s["status"] != "downloaded" for s in sk), f"zips/{t['folder']}.zip",
                   ", ".join(s["name"] for s in sk[:3])])
    for col, w in zip("ABCDEFG", (4, 30, 13, 12, 10, 40, 70)):
        ov.column_dimensions[col].width = w
    used = set()
    cols = ["Rank", "Skill", "Purpose", "How to use", "Source URL", "Repo", "Stars", "License",
            "Feedback evidence", "Trending", "Status", "Local folder", "Notes"]
    widths = [6, 28, 50, 50, 45, 28, 9, 13, 40, 9, 18, 40, 35]
    for t in topics:
        name = re.sub(r"[\[\]:*?/\\]", "", t["topic"])[:31]
        while name in used:
            name = name[:29] + str(len(used))
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.fill, c.font = hdr_fill, hdr_font
        for s in t["skills"]:
            ws.append([s["rank"], s["name"], s.get("purpose", ""), s.get("how_to_use", ""), s.get("url", ""),
                       s.get("repo", ""), s.get("stars"), s.get("license", ""), s.get("evidence", ""),
                       "yes" if s.get("trending") else "", s["status"], s.get("local", ""), s.get("notes", "")])
            r = ws.max_row
            if s.get("url"):
                ws.cell(r, 5).hyperlink = s["url"]
                ws.cell(r, 5).font = Font(color="0563C1", underline="single")
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(ROOT / "SkillGarden_Index.xlsx")

if __name__ == "__main__":
    main(download="--no-download" not in sys.argv, only=[a for a in sys.argv[1:] if not a.startswith("--")])
